"""The report: one HTML document (saved as .html, or printed to .pdf by the window), the
same sections as an Excel workbook (a sheet each), or one PowerPoint slide (the figure and
the tables on it, the other sections as text in its notes). No Qt.

Sections, in order; any can be left out. Images are named in the HTML (`<img src="figure.png">`)
and handed over apart, so the .html embeds them and the PDF printer loads them. Names are
the aliases (▤) where one is given.
"""

from __future__ import annotations

import datetime as _dt
import functools
import hashlib
import html
import io
import os
import statistics
from html.parser import HTMLParser
from pathlib import Path

from . import __version__, fcs, model, plots, record

SECTIONS = [("experiment", "Experiment"), ("figure", "Figure"), ("results", "Results"),
            ("replicates", "Replicates (runs pooled)"), ("samples", "Samples"),
            ("files", "Files"),
            ("procedure", "Procedure and reagents"),
            ("gating", "Gating"), ("instrument", "Acquisition and instrument"),
            ("checklist", "MIFlowCyt checklist")]


def num(v, nd: int = 0) -> str:
    """'70 543' (thin-space thousands), '' for nothing."""
    if v is None or v == "":
        return ""
    if isinstance(v, float) and nd:
        return f"{v:.{nd}f}"
    return f"{v:,.0f}".replace(",", " ")


@functools.lru_cache(maxsize=None)
def _sha256(path: str, mtime: float, size: int) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def file_info(path, base: str = "") -> list[str]:
    """[name, folder (under *base*), modified, SHA-256's first 12 characters]."""
    p = Path(path)
    try:
        st = p.stat()
    except OSError:
        return [p.name, str(p.parent), "not found", ""]
    folder = os.path.relpath(p.parent, base) if base else str(p.parent)
    return [p.name, "" if folder == "." else folder,
            f"{_dt.datetime.fromtimestamp(st.st_mtime):%Y-%m-%d %H:%M}",
            _sha256(str(p), st.st_mtime, st.st_size)[:12]]


def files_table(exp: model.Experiment) -> tuple[str, list[list[str]]]:
    """(the folder they share, a row per file): the FCS files, then the gates' file while
    the gates are still as in it."""
    paths = [s.path for s in exp.samples]
    gf = exp.gates_from()
    try:
        base = os.path.commonpath([str(Path(x).parent) for x in paths + [gf.get("path")]
                                   if x])
    except ValueError:                                     # two drives
        base = ""
    rows = [file_info(x, base) for x in paths]
    if gf:
        r = file_info(gf["path"], base)
        rows.append([f"{r[0]} — the gates, {gf['how']} (unchanged since)"] + r[1:])
    return base, rows


def summary_rows(exp: model.Experiment, rows: list[dict] | None = None) -> list[dict]:
    """The table of the figure: group, label, sub-population ('' for the one shown), its
    events, the chosen statistic, ÷ control (over the control's value in the population
    shown). The population each group's figure shows (plots.measured) and every one below
    it; unplotted samples left out."""
    rows = rows if rows is not None else exp.results()
    plotted = {(s.group, s.name) for s in exp.samples if s.plot}
    shown = {g: plots.measured(exp, g) for g in exp.groups()}
    ctrl = {}
    for g in exp.groups():
        c = exp.control(g)
        ctrl[g] = next((r[exp.stat] for r in rows if c is not None and r["sample"] == c.name
                        and r["group"] == g and r["population"] == shown[g]), None)
    out = []
    for r in rows:
        g, pop = r["group"], shown[r["group"]]
        if (g, r["sample"]) not in plotted:
            continue
        if r["population"] == pop or r["population"] in exp.below(g, pop):
            c, v = ctrl.get(g), r[exp.stat]
            out.append({"group": g, "label": r["label"],
                        "sub-population": "" if r["population"] == pop else r["population"],
                        "population": exp.name("population", r["population"]),
                        "count": r["count"], "% parent": r["% parent"], "value": v,
                        "÷ control": v / c if c and v is not None else None, "role": r["role"]})
    return out


def replicate_rows(exp: model.Experiment, sr: list[dict] | None = None) -> list[dict]:
    """The summary pooled over the groups shown under one name (runs of one condition): per
    label and sub-population, the runs' n, mean and SD of the value and of ÷ control (each
    run over its own control: comparable when the runs' intensities differ)."""
    sr = sr if sr is not None else summary_rows(exp)
    out, seen = [], set()
    for g in exp.groups():
        pool = exp.pool(g)
        name = exp.name("group", g)
        if len(pool) < 2 or name in seen:
            continue
        seen.add(name)
        keys = list(dict.fromkeys((exp.name("label", r["label"]), r["sub-population"])
                                  for r in sr if r["group"] in pool))
        main = " / ".join(dict.fromkeys(exp.name("population", plots.measured(exp, x))
                                        for x in pool))       # each run's: alike, mostly
        for lab, sub in keys:
            got = [r for r in sr if r["group"] in pool and r["sub-population"] == sub
                   and exp.name("label", r["label"]) == lab]
            v = [r["value"] for r in got if r["value"] is not None]
            q = [r["÷ control"] for r in got if r["÷ control"] is not None]
            out.append({"group": name, "label": lab,
                        "sub-population": exp.name("population", sub) if sub else "",
                        "population": exp.name("population", sub) if sub else main,
                        "n": len(v), "mean": statistics.mean(v) if v else None,
                        "sd": statistics.stdev(v) if len(v) > 1 else None,
                        "ratio mean": statistics.mean(q) if q else None,
                        "ratio sd": statistics.stdev(q) if len(q) > 1 else None,
                        "runs": ", ".join(f"{exp.run_of(r['group'])}: {num(r['value'])}"
                                          for r in got)})
    return out


def table(heads, rows, cls="") -> str:
    h = "".join(f"<th>{html.escape(str(x))}</th>" for x in heads)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>"
                   for r in rows)
    width = "" if {"kv", "files"} & set(cls.split()) else ' width="100%"'
    return (f'<table class="{cls}"{width} border="1" cellspacing="0" cellpadding="4">'
            f"<thead><tr>{h}</tr></thead><tbody>{body}</tbody></table>")


def _png(fig, dpi=200) -> bytes:
    b = io.BytesIO()
    fig.savefig(b, format="png", dpi=dpi, facecolor="white")
    return b.getvalue()


def build(exp: model.Experiment, sections=None, full_table: bool = False) -> tuple[str, dict]:
    """(html, {image name: png bytes})."""
    from matplotlib.figure import Figure

    sections = sections or [k for k, _ in SECTIONS]
    rec = record.clean(exp.record)
    g = rec["general"]
    acq = dict(record.acquisition(exp.samples))
    first = exp.samples[0].fcs if exp.samples else None
    date = g.get("date") or acq.get("date", "")
    title = g.get("title") or (Path(exp.path).stem if exp.path else "Flow cytometry")
    images, parts = {}, []
    N = lambda kind, x: exp.name(kind, x)  # noqa: E731
    parts.append(f"<h1>{html.escape(title)}</h1>")
    parts.append(f'<p class="meta">{html.escape(date)} · marker '
                 f"{html.escape(N('channel', exp.marker))} · "
                 f"Flowless {__version__} · report {_dt.datetime.now():%Y-%m-%d %H:%M}</p>")
    rows = exp.results()
    for key in sections:
        name = dict(SECTIONS)[key]
        if key == "experiment":
            items = [(k, g.get(k, "")) for k in record.GENERAL if g.get(k) and k != "title"]
            grp = [[exp.title(x)] + [rec["groups"].get(x, {}).get(f, "")
                                      for f in record.GROUP_FIELDS] for x in exp.groups()]
            parts.append(f"<h2>{name}</h2>")
            if items:
                parts.append(table(["", ""], items, "kv"))
            parts.append(table(["group"] + record.GROUP_FIELDS, grp))
        elif key == "figure":
            fig = plots.figure(exp, style={"background": "white"})
            images["figure.png"] = _png(fig, 250)
            w = int(plots.figure(exp, draw=False)[0] * 3.78)
            parts.append(f'<h2>{name}</h2><p><img src="figure.png" width="{min(w, 718)}"></p>'
                         f'<p class="note">{html.escape(_figure_legend(exp))}</p>')
        elif key == "results":
            parts.append(f"<h2>{name}</h2>")
            sr = summary_rows(exp, rows)
            parts.append(table(["group", "label", "population", "events", "% parent",
                                exp.stat, "÷ control"],
                               [[exp.title(r["group"]), N("label", r["label"]),
                                 r["population"], num(r["count"]),
                                 num(r["% parent"], 1), num(r["value"]),
                                 num(r["÷ control"], 2)] for r in sr]))
            if full_table:
                parts.append("<h3>Every sample and population</h3>")
                parts.append(table(
                    ["group", "label", "population", "count", "% parent", "geometric mean",
                     "median", "mean", "CV %", "÷ control"],
                    [[exp.title(r["group"]), N("label", r["label"]),
                      N("population", r["population"]) + (" ◆" if r["own gate"] else ""),
                      num(r["count"]), num(r["% parent"], 1), num(r["geometric mean"]),
                      num(r["median"]), num(r["mean"]), num(r["CV %"], 1),
                      num(r["÷ control"], 2)] for r in rows], "small"))
                if any(r["own gate"] for r in rows):
                    parts.append('<p class="note">◆ the sample\'s own version of the gate.</p>')
            parts.append(f'<p class="note">{html.escape(stat_note(exp))}</p>')
        elif key == "replicates":
            rr = replicate_rows(exp, summary_rows(exp, rows))
            parts.append(f"<h2>{name}</h2>")
            if rr:
                parts.append(table(["group", "label", "population", "runs",
                                    f"{exp.stat} (mean ± SD)", "÷ control (mean ± SD)",
                                    "each run"],
                                   [[r["group"], r["label"], r["population"],
                                     r["n"], f"{num(r['mean'])} ± {num(r['sd'])}",
                                     f"{num(r['ratio mean'], 2)} ± {num(r['ratio sd'], 2)}",
                                     r["runs"]] for r in rr], "small"))
                parts.append('<p class="note">Groups shown under one name (▤) are runs of '
                             "one condition. SD over the runs (n − 1); ÷ control: each run "
                             "over its own control.</p>")
            else:
                parts.append('<p class="note">No two groups are shown under one name: give '
                             "two runs of a condition the same name (▤ Aliases) to pool "
                             "them.</p>")
        elif key == "samples":
            extra = [k for k in dict.fromkeys(k for s in exp.samples for k in s.fields)]
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(
                ["group", "label", "role", "file", "well", "acquired", "events", "µL"] + extra,
                [[exp.title(s.group), N("label", s.label), s.role, Path(s.path).name,
                  fcs.pick(s.fcs.meta, ("$WELLID",)), fcs.when(s.fcs.meta),
                  num(s.fcs.n), num(fcs.volume_ul(s.fcs.meta), 1)]
                 + [s.fields.get(k, "") for k in extra] for s in exp.samples], "small"))
        elif key == "files":
            base, frows = files_table(exp)
            parts.append(f"<h2>{name}</h2>")
            parts.append(f'<p class="note">{html.escape(base and "in " + base + " · ")}'
                         "SHA-256: its first 12 characters</p>")
            parts.append(table(["file", "folder", "modified", "SHA-256"], frows, "small files"))
        elif key == "procedure":
            parts.append(f"<h2>{name}</h2>")
            if rec["procedure"]:
                parts.append("<ol>" + "".join(f"<li>{html.escape(x)}</li>"
                                              for x in rec["procedure"]) + "</ol>")
            if rec["reagents"]:
                cols = [c for c in record.REAGENT_FIELDS
                        if any(r.get(c) for r in rec["reagents"])]
                parts.append(table(cols + ["reads as"],
                                   [[r.get(c, "") for c in cols] + [record.describe(r)]
                                    for r in rec["reagents"]]))
        elif key == "gating":
            parts.append(f"<h2>{name}</h2>")
            for k, grp in enumerate(exp.groups()):
                gates = exp.gates.get(grp, [])
                if not gates:
                    parts.append(f"<p><b>{html.escape(exp.title(grp))}</b>: no gate (all "
                                 "events)</p>")
                    continue
                desc = [[N("population", x.name), N("population", x.parent or model.ALL),
                         f"{x.x} × {x.y} polygon, {len(x.vertices)} vertices"
                         if x.kind == "polygon" else
                         f"{x.x} from {num(x.lo)} to {num(x.hi) or 'the top'}",
                         ", ".join(N("label", s.label) for s in exp.in_group(grp)
                                   if x.name in s.gates)] for x in gates]
                polys = [x for x in gates if x.kind == "polygon"]
                parts.append(f"<h3>{html.escape(exp.title(grp))}</h3>")
                if polys:
                    fig = Figure(figsize=(2.6 * len(polys), 2.6))
                    for ax, x in zip(fig.subplots(1, len(polys), squeeze=False)[0], polys):
                        plots.dot_axes(ax, exp, exp.in_group(grp), x.x, x.y, x.parent, [x],
                                       {"label_pt": 7, "xtick_pt": 6, "ytick_pt": 6,
                                        "gate_pt": 6.5})
                    fig.tight_layout()
                    images[f"gate{k}.png"] = _png(fig, 150)
                    parts.append(f'<p><img src="gate{k}.png" width="{int(250 * len(polys))}">'
                                 "</p>")
                parts.append(table(["population", "parent", "gate", "own version in"], desc))
        elif key == "instrument":
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(["", ""], record.acquisition(exp.samples) +
                               record.instrument(first), "kv"))
        elif key == "checklist":
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(["item", "", "from"],
                               [[i, "✓" if ok else "✗ missing", w] for i, ok, w in
                                record.miflowcyt(rec, exp.samples, exp.gates)]))
    return _page(title, "\n".join(parts)), images


Y_TEXT = {"mode": "% of each histogram’s maximum (normalised to mode)",
          "area": "% of the population’s events per channel (each histogram’s area 100 %)"}


def _figure_legend(exp) -> str:
    o = plots.options(exp)
    hist = [p for p in plots.panels(exp) if p.get("kind") == "histogram"]
    pops = {p.get("population") or exp.main_population(p.get("group", "")) for p in hist}
    chans = {exp.name("channel", p.get("x") or exp.marker) for p in hist}
    if not hist:
        return ""
    return (f"Histograms of {', '.join(sorted(chans))} on {', '.join(sorted(pops))}; "
            f"{Y_TEXT.get(o['y'], 'events per channel')}"
            f" (256 channels over the log scale{', smoothed' if o['smooth'] else ''}).")


def stat_note(exp) -> str:
    shown = ", ".join(dict.fromkeys(exp.name("population", plots.measured(exp, g))
                                    for g in exp.groups()))
    return (f"{exp.stat.capitalize()} of {exp.name('channel', exp.marker)} in each population "
            "(geometric mean: of the positive values); the summary: the population each "
            f"group's figure shows ({shown}) and those below it. ÷ control: over the group's "
            "negative control (else secondary-only, else unstained) in that population.")


def _page(title, body) -> str:
    return f"""<!doctype html><html><head><meta charset="utf-8">
<title>{html.escape(title)}</title><style>
body {{ font-family: Arial, sans-serif; font-size: 10pt; color: #0b0b0b; }}
@media screen {{ body {{ max-width: 190mm; margin: 12mm auto; }} }}
h1 {{ font-size: 16pt; margin-bottom: 2pt; }} h2 {{ font-size: 12.5pt; margin-top: 16pt;
     border-bottom: 1px solid #c3c2b7; }} h3 {{ font-size: 10.5pt; }}
.meta, .note {{ color: #52514e; font-size: 9pt; }}
table {{ border-collapse: collapse; font-size: 9pt; }}
th, td {{ border: 1px solid #c3c2b7; padding: 3px 6px; text-align: left; }}
th {{ background: #f0efec; }} table.kv th {{ display: none; }}
table.small {{ font-size: 7.5pt; }} table.files td {{ white-space: nowrap; }}
</style></head><body>
{body}
</body></html>"""


def save_html(html_text: str, images: dict, path) -> None:
    import base64
    for k, v in images.items():
        html_text = html_text.replace(f'src="{k}"',
                                      f'src="data:image/png;base64,{base64.b64encode(v).decode()}"')
    Path(path).write_text(html_text, encoding="utf-8")


def save_xlsx(exp: model.Experiment, path, sections=None) -> None:
    """A sheet (or two) per section, in the sections' order; numbers as numbers. None: every
    section but the figure."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    sections = sections or [k for k, _ in SECTIONS if k != "figure"]
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    N = lambda kind, x: exp.name(kind, x)  # noqa: E731

    def sheet(title, heads, rows, ws=None):
        ws = ws or wb.create_sheet(title)
        ws.append(heads)
        for c in ws[ws.max_row]:
            c.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                40, max(8, *(len(str(c.value or "")) + 1 for c in col)))
        return ws

    rows = exp.results()
    rec = record.clean(exp.record)
    for key in sections:
        if key == "results":
            sheet("Summary", ["group", "label", "role", "population", "events", "% parent",
                              exp.stat, "÷ control"],
                  [[exp.title(r["group"]), N("label", r["label"]), r["role"],
                    r["population"], r["count"], r["% parent"],
                    r["value"], r["÷ control"]] for r in summary_rows(exp, rows)])
            heads = ["group", "label", "role", "sample", "population", "parent", "count",
                     "% parent", "geometric mean", "median", "mean", "CV %", "÷ control",
                     "own gate"]
            sheet("Results", heads, [[exp.title(r["group"]) if h == "group" else
                                      N(h, r[h]) if h in ("label", "sample", "population")
                                      else r[h] for h in heads] for r in rows])
        elif key == "replicates":
            rr = replicate_rows(exp, summary_rows(exp, rows))
            if rr:
                sheet("Replicates", ["group", "label", "population", "runs",
                                     f"{exp.stat} mean", "SD", "÷ control mean",
                                     "÷ control SD", "each run"],
                      [[r["group"], r["label"], r["population"], r["n"], r["mean"], r["sd"],
                        r["ratio mean"], r["ratio sd"], r["runs"]] for r in rr])
        elif key == "replicates":
            rr = replicate_rows(exp, summary_rows(exp, rows))
            parts.append(f"<h2>{name}</h2>")
            if rr:
                parts.append(table(["group", "label", "population", "runs",
                                    f"{exp.stat} (mean ± SD)", "÷ control (mean ± SD)",
                                    "each run"],
                                   [[r["group"], r["label"], r["population"],
                                     r["n"], f"{num(r['mean'])} ± {num(r['sd'])}",
                                     f"{num(r['ratio mean'], 2)} ± {num(r['ratio sd'], 2)}",
                                     r["runs"]] for r in rr], "small"))
                parts.append('<p class="note">Groups shown under one name (▤) are runs of '
                             "one condition. SD over the runs (n − 1); ÷ control: each run "
                             "over its own control.</p>")
            else:
                parts.append('<p class="note">No two groups are shown under one name: give '
                             "two runs of a condition the same name (▤ Aliases) to pool "
                             "them.</p>")
        elif key == "samples":
            extra = list(dict.fromkeys(k for s in exp.samples for k in s.fields))
            sheet("Samples", ["group", "label", "role", "plot", "file", "well", "date", "start",
                              "events", "volume µL", "acquisition s"] + extra,
                  [[exp.title(s.group), N("label", s.label), s.role, s.plot, Path(s.path).name,
                    fcs.pick(s.fcs.meta, ("$WELLID",)), fcs.pick(s.fcs.meta, ("$DATE",)),
                    fcs.pick(s.fcs.meta, ("$BTIM",)), s.fcs.n, fcs.volume_ul(s.fcs.meta),
                    fcs.acquired_s(s.fcs.meta)]
                   + [s.fields.get(k, "") for k in extra] for s in exp.samples])
        elif key == "files":
            sheet("Files", ["file", "folder", "modified", "SHA-256 (12)"],
                  [file_info(s.path) for s in exp.samples] +
                  ([[f"{Path(gf['path']).name} (the gates, {gf['how']})",
                     *file_info(gf["path"])[1:]]] if (gf := exp.gates_from()) else []))
        elif key == "experiment":
            ws = sheet("Experiment", ["field", "value"],
                       [[k, rec["general"].get(k, "")] for k in record.GENERAL])
            ws.append([])
            ws.append(["group"] + record.GROUP_FIELDS)
            for g in exp.groups():
                ws.append([exp.title(g)] + [rec["groups"].get(g, {}).get(f, "")
                                            for f in record.GROUP_FIELDS])
        elif key == "procedure":
            ws = sheet("Procedure", record.REAGENT_FIELDS + ["reads as"],
                       [[r.get(f, "") for f in record.REAGENT_FIELDS] + [record.describe(r)]
                        for r in rec["reagents"]])
            ws.append([])
            ws.append(["step", "procedure"])
            for i, st in enumerate(rec["procedure"], 1):
                ws.append([i, st])
        elif key == "gating":
            gate_rows = [[exp.title(g), "(group)", x.name, x.parent or model.ALL, x.kind, x.x,
                          x.y, x.lo, x.hi, "; ".join(f"{a:.0f},{b:.0f}" for a, b in x.vertices)]
                         for g, gs in exp.gates.items() for x in gs]
            gate_rows += [[exp.title(s.group), s.name, x.name, x.parent or model.ALL, x.kind,
                           x.x, x.y, x.lo, x.hi,
                           "; ".join(f"{a:.0f},{b:.0f}" for a, b in x.vertices)]
                          for s in exp.samples for x in s.gates.values()]
            sheet("Gates", ["group", "sample", "population", "parent", "kind", "x", "y", "lo",
                            "hi", "vertices"], gate_rows)
        elif key == "instrument":
            sheet("Instrument", ["", ""], record.acquisition(exp.samples) +
                  record.instrument(exp.samples[0].fcs if exp.samples else None))
            keys = list(dict.fromkeys(k for s in exp.samples for k in s.fcs.meta))
            sheet("Keywords", ["keyword"] + [s.name for s in exp.samples],
                  [[k] + [s.fcs.meta.get(k, "") for s in exp.samples] for k in keys])
        elif key == "checklist":
            sheet("MIFlowCyt", ["item", "", "from"],
                  [[i, "✓" if ok else "✗ missing", w]
                   for i, ok, w in record.miflowcyt(rec, exp.samples, exp.gates)])
        elif key == "figure":
            from openpyxl.drawing.image import Image
            ws = wb.create_sheet("Figure")
            img = Image(io.BytesIO(_png(plots.figure(exp, style={"background": "white"}), 150)))
            ws.add_image(img, "A1")
    if not wb.sheetnames:
        wb.create_sheet("Summary")
    wb.save(path)


# --- PowerPoint: one slide (the figure, the tables), the rest in its notes ----------------

class _Text(HTMLParser):
    """The report's HTML as plain text: headings on their own line, a table row per line
    (cells joined by ' | '), list items as '- '."""

    def __init__(self):
        super().__init__()
        self.out, self.row, self.cell, self.skip, self.head = [], None, None, False, False

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "title"):
            self.skip = True
        elif tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            self.cell = []
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in ("h1", "h2", "h3"):
            self.out.append("\n\n")
            self.head = True
        elif tag in ("p", "br", "table", "ol", "ul"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("style", "title"):
            self.skip = False
        elif tag in ("td", "th") and self.row is not None:
            self.row.append("".join(self.cell or []).strip())
            self.cell = None
        elif tag == "tr":
            if any(self.row or []):
                self.out.append("\n" + " | ".join(self.row))
            self.row = None
        elif tag in ("h1", "h2", "h3"):
            self.head = False
            self.out.append("\n")

    def handle_data(self, d):
        if self.skip:
            return
        d = d.upper() if self.head else d
        (self.cell if self.cell is not None else self.out).append(d)

    def text(self) -> str:
        t = "".join(self.out)
        return "\n".join(x.rstrip() for x in t.splitlines()).replace("\n\n\n", "\n\n").strip()


def as_text(html_text: str) -> str:
    p = _Text()
    p.feed(html_text)
    return p.text()


SLIDE = ("figure", "results", "replicates")      # on the slide; the other sections: its notes


def where(ext: str, key: str) -> str:
    """Where a section goes in that format: 'slide' / 'notes' (.pptx), 'sheet' (.xlsx), ''."""
    if ext == ".pptx":
        return "slide" if key in SLIDE else "notes"
    return "sheet" if ext == ".xlsx" else ""


def pptx_parts(exp, sections, full_table=False) -> tuple[str, dict, str]:
    """(the slide's html, its images, the notes as text): the preview of the slide."""
    on = [k for k in sections if k in SLIDE]
    text, images = build(exp, on, full_table)
    notes, _ = build(exp, [k for k in sections if k not in SLIDE])
    return text, images, as_text(notes)


PAD = 0.2                                        # cm: a cell's left + right margins


def _text_cm(t: str, pt: float, bold=False) -> float:
    """About how wide *t* is in Arial (0.55 em a character; bold 0.6)."""
    return len(t) * pt * 0.0353 * (0.6 if bold else 0.55)


def _lay(heads, body, width: float, pt: float) -> tuple[list[float], list[float]]:
    """Column widths (cm) from the contents, as wide as *width*: when they don't fit, the
    widest text column gives way and wraps; and each row's height from its wrapping."""
    text = [max([_text_cm(str(r[c]), pt) for r in body] or [0]) + PAD for c in range(len(heads))]
    need = [max(t, _text_cm(str(h), pt, True) + PAD) for t, h in zip(text, heads)]
    if sum(need) > width:                  # the widest body text gives way, then the headers
        c = max(range(len(text)), key=lambda c: text[c])
        need[c] = max(width - (sum(need) - need[c]), 2.0)
        if sum(need) > width:
            need = [max(t, 1.2) for t in text]
            need = [w * width / sum(need) for w in need] if sum(need) > width else                 [w + (width - sum(need)) / len(need) for w in need]
    else:
        need = [w + (width - sum(need)) / len(need) for w in need]
    line = pt * 0.0353 * 1.2
    heights = [max(0.5, 0.12 + line * max(
        -(-_text_cm(str(v), pt, i == 0) // max(need[c] - PAD, 0.1)) for c, v in enumerate(vals)))
        for i, vals in enumerate([heads] + body)]
    return need, heights


def save_pptx(exp: model.Experiment, path, sections=None) -> None:
    """One slide, a figure page like the bench's: the title, the figure at its size, the
    summary (and replicates) as native tables below it; every other section ticked as
    text in the slide's notes."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Cm, Pt

    sections = sections or [k for k, _ in SECTIONS]
    rec = record.clean(exp.record)
    g = rec["general"]
    acq = dict(record.acquisition(exp.samples))
    title = g.get("title") or (Path(exp.path).stem if exp.path else "Flow cytometry")
    rows = exp.results()
    sr = summary_rows(exp, rows)
    rr = replicate_rows(exp, sr) if "replicates" in sections else []
    tables = []
    if "results" in sections:
        tables.append((["group", "label", "population", "events", "% parent", exp.stat,
                        "÷ control"],
                       [[exp.title(r["group"]), exp.name("label", r["label"]),
                         r["population"], num(r["count"]),
                         num(r["% parent"], 1), num(r["value"]), num(r["÷ control"], 2)]
                        for r in sr]))
    if rr:
        tables.append((["group", "label", "population", "runs", f"{exp.stat} (mean ± SD)",
                        "÷ control (mean ± SD)"],
                       [[r["group"], r["label"], r["population"], str(r["n"]),
                         f"{num(r['mean'])} ± {num(r['sd'])}",
                         f"{num(r['ratio mean'], 2)} ± {num(r['ratio sd'], 2)}"] for r in rr]))
    m, head_h, pt = 0.6, 0.9, 9                            # cm, cm, table font
    fw = fh = 0.0
    if "figure" in sections:
        fw, fh = (v / 10 for v in plots.figure(exp, draw=False))
    W = max(16.0, fw + 2 * m)
    laid = [_lay(heads, body, W - 2 * m, pt) for heads, body in tables]
    H = m + head_h + (fh + 0.3 if fw else 0) + sum(sum(hs) + 0.4 for _, hs in laid) + m
    pres = Presentation()
    pres.slide_width, pres.slide_height = Cm(W), Cm(max(H, 9.0))
    pres.core_properties.title = title
    slide = pres.slides.add_slide(pres.slide_layouts[6])            # blank
    tb = slide.shapes.add_textbox(Cm(m), Cm(m), Cm(W - 2 * m), Cm(head_h))
    tb.name = "Title"
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    r = tf.paragraphs[0].add_run()
    r.text = title
    r.font.size, r.font.bold, r.font.name = Pt(14), True, "Arial"
    r = tf.paragraphs[0].add_run()
    r.text = f"   {g.get('date') or acq.get('date', '')} · {exp.name('channel', exp.marker)}"
    r.font.size, r.font.name = Pt(9), "Arial"
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    y = m + head_h
    if fw:
        png = _png(plots.figure(exp, style={"background": "white"}), 300)
        pic = slide.shapes.add_picture(io.BytesIO(png), Cm((W - fw) / 2), Cm(y), Cm(fw), Cm(fh))
        pic.name = "Figure"
        y += fh + 0.3
    for k, ((heads, body), (widths, heights)) in enumerate(zip(tables, laid)):
        shape = slide.shapes.add_table(len(body) + 1, len(heads), Cm(m), Cm(y),
                                       Cm(sum(widths)), Cm(sum(heights)))
        shape.name = ["Summary", "Replicates"][k] if len(tables) > 1 else "Summary"
        tbl = shape.table
        tbl.first_row = True
        tbl.horz_banding = False
        shape._element.graphic.graphicData.tbl[0][-1].text = \
            "{5940675A-B579-460E-94D1-54222C63F5DA}"            # no style, table grid
        for c, w in enumerate(widths):
            tbl.columns[c].width = Cm(w)
        for i, vals in enumerate([heads] + body):
            tbl.rows[i].height = Cm(heights[i])
            for c, v in enumerate(vals):
                cell = tbl.cell(i, c)
                cell.margin_top = cell.margin_bottom = Cm(0.03)
                cell.margin_left = cell.margin_right = Cm(PAD / 2)
                cell.text = str(v)
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if c >= 3 and i else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.size, run.font.name, run.font.bold = Pt(pt), "Arial", i == 0
        y += sum(heights) + 0.4
    notes = [k for k in sections if k not in SLIDE]
    if notes:
        text, _ = build(exp, notes)
        slide.notes_slide.notes_text_frame.text = as_text(text)
    pres.save(path)
