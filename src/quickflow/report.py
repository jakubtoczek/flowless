"""The report: one HTML document (saved as .html, or printed to .pdf by the window), and
the same content as an Excel workbook. No Qt.

Sections, in order; any can be left out. Images are named in the HTML (`<img src="figure.png">`)
and handed over apart, so the .html embeds them and the PDF printer loads them. Names are
the aliases (▤) where one is given.
"""

from __future__ import annotations

import datetime as _dt
import html
import io
from pathlib import Path

from . import __version__, fcs, model, plots, record

SECTIONS = [("experiment", "Experiment"), ("figure", "Figure"), ("results", "Results"),
            ("samples", "Samples"), ("procedure", "Procedure and reagents"),
            ("gating", "Gating"), ("instrument", "Acquisition and instrument"),
            ("checklist", "MIFlowCyt checklist")]


def num(v, nd: int = 0) -> str:
    """'70 543' (thin-space thousands), '' for nothing."""
    if v is None or v == "":
        return ""
    if isinstance(v, float) and nd:
        return f"{v:.{nd}f}"
    return f"{v:,.0f}".replace(",", " ")


def summary_rows(exp: model.Experiment, rows: list[dict] | None = None) -> list[dict]:
    """The table of the figure: group, label, sub-population ('' for the main one), its
    events, the chosen statistic, ÷ control. The main population and every one below it;
    unplotted samples left out."""
    rows = rows if rows is not None else exp.results()
    plotted = {(s.group, s.name) for s in exp.samples if s.plot}
    out = []
    for r in rows:
        if (r["group"], r["sample"]) not in plotted:
            continue
        main = exp.main_population(r["group"])
        if r["population"] == main or r["population"] in exp.below(r["group"], main):
            out.append({"group": r["group"], "label": r["label"],
                        "sub-population": "" if r["population"] == main else r["population"],
                        "count": r["count"], "value": r[exp.stat],
                        "÷ control": r["÷ control"], "role": r["role"]})
    return out


def table(heads, rows, cls="") -> str:
    h = "".join(f"<th>{html.escape(str(x))}</th>" for x in heads)
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>"
                   for r in rows)
    width = "" if cls == "kv" else ' width="100%"'
    return (f'<table class="{cls}"{width} border="1" cellspacing="0" cellpadding="4">'
            f"<thead><tr>{h}</tr></thead><tbody>{body}</tbody></table>")


def _png(fig, dpi=200) -> bytes:
    b = io.BytesIO()
    fig.savefig(b, format="png", dpi=dpi, facecolor="white")
    return b.getvalue()


def build(exp: model.Experiment, sections=None, full_table: bool = False) -> tuple[str, dict]:
    """(html, {image name: png bytes})."""
    from matplotlib.figure import Figure

    sections = sections or [k for k, _ in SECTIONS]
    rec = record.clean(exp.record)
    g = rec["general"]
    acq = dict(record.acquisition(exp.samples))
    first = exp.samples[0].fcs if exp.samples else None
    date = g.get("date") or acq.get("date", "")
    title = g.get("title") or (Path(exp.path).stem if exp.path else "Flow cytometry")
    images, parts = {}, []
    N = lambda kind, x: exp.name(kind, x)  # noqa: E731
    parts.append(f"<h1>{html.escape(title)}</h1>")
    parts.append(f'<p class="meta">{html.escape(date)} · marker '
                 f"{html.escape(N('channel', exp.marker))} · "
                 f"QuickFlow {__version__} · report {_dt.datetime.now():%Y-%m-%d %H:%M}</p>")
    rows = exp.results()
    for key in sections:
        name = dict(SECTIONS)[key]
        if key == "experiment":
            items = [(k, g.get(k, "")) for k in record.GENERAL if g.get(k) and k != "title"]
            grp = [[N("group", x)] + [rec["groups"].get(x, {}).get(f, "")
                                      for f in record.GROUP_FIELDS] for x in exp.groups()]
            parts.append(f"<h2>{name}</h2>")
            if items:
                parts.append(table(["", ""], items, "kv"))
            parts.append(table(["group"] + record.GROUP_FIELDS, grp))
        elif key == "figure":
            fig = plots.figure(exp, style={"background": "white"})
            images["figure.png"] = _png(fig, 250)
            w = int(plots.figure(exp, draw=False)[0] * 3.78)
            parts.append(f'<h2>{name}</h2><p><img src="figure.png" width="{min(w, 718)}"></p>'
                         f'<p class="note">{html.escape(_figure_legend(exp))}</p>')
        elif key == "results":
            parts.append(f"<h2>{name}</h2>")
            sr = summary_rows(exp, rows)
            parts.append(table(["group", "label", "sub-population", "events", exp.stat,
                                "÷ control"],
                               [[N("group", r["group"]), N("label", r["label"]),
                                 N("population", r["sub-population"]) or "—", num(r["count"]),
                                 num(r["value"]), num(r["÷ control"], 2)] for r in sr]))
            if full_table:
                parts.append("<h3>Every sample and population</h3>")
                parts.append(table(
                    ["group", "label", "population", "count", "% parent", "geometric mean",
                     "median", "mean", "CV %", "÷ control"],
                    [[N("group", r["group"]), N("label", r["label"]),
                      N("population", r["population"]) + (" ◆" if r["own gate"] else ""),
                      num(r["count"]), num(r["% parent"], 1), num(r["geometric mean"]),
                      num(r["median"]), num(r["mean"]), num(r["CV %"], 1),
                      num(r["÷ control"], 2)] for r in rows], "small"))
                if any(r["own gate"] for r in rows):
                    parts.append('<p class="note">◆ the sample\'s own version of the gate.</p>')
            parts.append(f'<p class="note">{html.escape(stat_note(exp))}</p>')
        elif key == "samples":
            extra = [k for k in dict.fromkeys(k for s in exp.samples for k in s.fields)]
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(
                ["group", "label", "role", "file", "well", "acquired", "events", "µL"] + extra,
                [[N("group", s.group), N("label", s.label), s.role, Path(s.path).name,
                  fcs.pick(s.fcs.meta, ("$WELLID",)), fcs.when(s.fcs.meta),
                  num(s.fcs.n), num(fcs.volume_ul(s.fcs.meta), 1)]
                 + [s.fields.get(k, "") for k in extra] for s in exp.samples], "small"))
        elif key == "procedure":
            parts.append(f"<h2>{name}</h2>")
            if rec["procedure"]:
                parts.append("<ol>" + "".join(f"<li>{html.escape(x)}</li>"
                                              for x in rec["procedure"]) + "</ol>")
            if rec["reagents"]:
                cols = [c for c in record.REAGENT_FIELDS
                        if any(r.get(c) for r in rec["reagents"])]
                parts.append(table(cols + ["reads as"],
                                   [[r.get(c, "") for c in cols] + [record.describe(r)]
                                    for r in rec["reagents"]]))
        elif key == "gating":
            parts.append(f"<h2>{name}</h2>")
            for k, grp in enumerate(exp.groups()):
                gates = exp.gates.get(grp, [])
                if not gates:
                    parts.append(f"<p><b>{html.escape(N('group', grp))}</b>: no gate (all "
                                 "events)</p>")
                    continue
                desc = [[N("population", x.name), N("population", x.parent or model.ALL),
                         f"{x.x} × {x.y} polygon, {len(x.vertices)} vertices"
                         if x.kind == "polygon" else
                         f"{x.x} from {num(x.lo)} to {num(x.hi) or 'the top'}",
                         ", ".join(N("label", s.label) for s in exp.in_group(grp)
                                   if x.name in s.gates)] for x in gates]
                polys = [x for x in gates if x.kind == "polygon"]
                parts.append(f"<h3>{html.escape(N('group', grp))}</h3>")
                if polys:
                    fig = Figure(figsize=(2.6 * len(polys), 2.6))
                    for ax, x in zip(fig.subplots(1, len(polys), squeeze=False)[0], polys):
                        plots.dot_axes(ax, exp, exp.in_group(grp), x.x, x.y, x.parent, [x],
                                       {"label_pt": 7, "tick_pt": 6, "gate_pt": 6.5})
                    fig.tight_layout()
                    images[f"gate{k}.png"] = _png(fig, 150)
                    parts.append(f'<p><img src="gate{k}.png" width="{int(250 * len(polys))}">'
                                 "</p>")
                parts.append(table(["population", "parent", "gate", "own version in"], desc))
        elif key == "instrument":
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(["", ""], record.acquisition(exp.samples) +
                               record.instrument(first), "kv"))
        elif key == "checklist":
            parts.append(f"<h2>{name}</h2>")
            parts.append(table(["item", "", "from"],
                               [[i, "✓" if ok else "✗ missing", w] for i, ok, w in
                                record.miflowcyt(rec, exp.samples, exp.gates)]))
    return _page(title, "\n".join(parts)), images


def _figure_legend(exp) -> str:
    o = plots.options(exp)
    hist = [p for p in plots.panels(exp) if p.get("kind") == "histogram"]
    pops = {p.get("population") or exp.main_population(p.get("group", "")) for p in hist}
    chans = {exp.name("channel", p.get("x") or exp.marker) for p in hist}
    if not hist:
        return ""
    return (f"Histograms of {', '.join(sorted(chans))} on {', '.join(sorted(pops))}; "
            f"{'% of each histogram’s maximum (normalised to mode)' if o['y'] == 'mode' else 'events per channel'}"
            f" (256 channels over the log scale{', smoothed' if o['smooth'] else ''}).")


def stat_note(exp) -> str:
    return (f"{exp.stat.capitalize()} of {exp.name('channel', exp.marker)} in each population "
            "(geometric mean: of the positive values). ÷ control: over the group's "
            "negative control (else secondary-only, else unstained), main population.")


def _page(title, body) -> str:
    return f"""<!doctype html><html><head><meta charset="utf-8">
<title>{html.escape(title)}</title><style>
body {{ font-family: Arial, sans-serif; font-size: 10pt; color: #0b0b0b; }}
@media screen {{ body {{ max-width: 190mm; margin: 12mm auto; }} }}
h1 {{ font-size: 16pt; margin-bottom: 2pt; }} h2 {{ font-size: 12.5pt; margin-top: 16pt;
     border-bottom: 1px solid #c3c2b7; }} h3 {{ font-size: 10.5pt; }}
.meta, .note {{ color: #52514e; font-size: 9pt; }}
table {{ border-collapse: collapse; font-size: 9pt; }}
th, td {{ border: 1px solid #c3c2b7; padding: 3px 6px; text-align: left; }}
th {{ background: #f0efec; }} table.kv th {{ display: none; }}
table.small {{ font-size: 7.5pt; }}
</style></head><body>
{body}
</body></html>"""


def save_html(html_text: str, images: dict, path) -> None:
    import base64
    for k, v in images.items():
        html_text = html_text.replace(f'src="{k}"',
                                      f'src="data:image/png;base64,{base64.b64encode(v).decode()}"')
    Path(path).write_text(html_text, encoding="utf-8")


def save_xlsx(exp: model.Experiment, path) -> None:
    """Sheets: Summary, Results, Samples, Experiment, Gates, Keywords. Numbers as numbers."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)
    N = lambda kind, x: exp.name(kind, x)  # noqa: E731

    def sheet(title, heads, rows, ws=None):
        ws = ws or wb.create_sheet(title)
        ws.append(heads)
        for c in ws[ws.max_row]:
            c.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                40, max(8, *(len(str(c.value or "")) + 1 for c in col)))
        return ws

    rows = exp.results()
    ws = wb.active
    ws.title = "Summary"
    sheet("Summary", ["group", "label", "role", "sub-population", "events", exp.stat,
                      "÷ control"],
          [[N("group", r["group"]), N("label", r["label"]), r["role"],
            N("population", r["sub-population"]), r["count"], r["value"], r["÷ control"]]
           for r in summary_rows(exp, rows)], ws)
    heads = ["group", "label", "role", "sample", "population", "parent", "count", "% parent",
             "geometric mean", "median", "mean", "CV %", "÷ control", "own gate"]
    sheet("Results", heads, [[N(h, r[h]) if h in ("group", "label", "sample", "population")
                              else r[h] for h in heads] for r in rows])
    extra = list(dict.fromkeys(k for s in exp.samples for k in s.fields))
    sheet("Samples", ["group", "label", "role", "plot", "file", "folder", "well", "date", "start",
                      "events", "volume µL", "acquisition s"] + extra,
          [[N("group", s.group), N("label", s.label), s.role, s.plot, Path(s.path).name,
            str(Path(s.path).parent), fcs.pick(s.fcs.meta, ("$WELLID",)),
            fcs.pick(s.fcs.meta, ("$DATE",)), fcs.pick(s.fcs.meta, ("$BTIM",)), s.fcs.n,
            fcs.volume_ul(s.fcs.meta), fcs.acquired_s(s.fcs.meta)]
           + [s.fields.get(k, "") for k in extra] for s in exp.samples])
    rec = record.clean(exp.record)
    ws = sheet("Experiment", ["field", "value"],
               [[k, rec["general"].get(k, "")] for k in record.GENERAL])
    ws.append([])
    ws.append(["group"] + record.GROUP_FIELDS)
    for g in exp.groups():
        ws.append([N("group", g)] + [rec["groups"].get(g, {}).get(f, "")
                                     for f in record.GROUP_FIELDS])
    ws.append([])
    ws.append(record.REAGENT_FIELDS + ["reads as"])
    for r in rec["reagents"]:
        ws.append([r.get(f, "") for f in record.REAGENT_FIELDS] + [record.describe(r)])
    ws.append([])
    ws.append(["procedure"])
    for i, st in enumerate(rec["procedure"], 1):
        ws.append([i, st])
    ws.append([])
    for k, v in record.acquisition(exp.samples) + \
            record.instrument(exp.samples[0].fcs if exp.samples else None):
        ws.append([k, v])
    gate_rows = [[g, "(group)", x.name, x.parent or model.ALL, x.kind, x.x, x.y, x.lo, x.hi,
                  "; ".join(f"{a:.0f},{b:.0f}" for a, b in x.vertices)]
                 for g, gs in exp.gates.items() for x in gs]
    gate_rows += [[s.group, s.name, x.name, x.parent or model.ALL, x.kind, x.x, x.y, x.lo, x.hi,
                   "; ".join(f"{a:.0f},{b:.0f}" for a, b in x.vertices)]
                  for s in exp.samples for x in s.gates.values()]
    sheet("Gates", ["group", "sample", "population", "parent", "kind", "x", "y", "lo", "hi",
                    "vertices"], gate_rows)
    keys = list(dict.fromkeys(k for s in exp.samples for k in s.fcs.meta))
    sheet("Keywords", ["keyword"] + [s.name for s in exp.samples],
          [[k] + [s.fcs.meta.get(k, "") for s in exp.samples] for k in keys])
    wb.save(path)
