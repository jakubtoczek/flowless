"""The whole window, driven without a screen: add a run, import its procedure sheet and
workspace gates, every tab drawn, a gate drawn and edited with the mouse, a sample's own
gate, the figure's panels, aliases, figure / values / report saved, save and reopen, undo.

    .venv/Scripts/python misc/smoke_test.py [DATA] [--shots DIR]

DATA: the example data folder (default ../data/data). --shots: a screenshot per tab there.
"""

from __future__ import annotations

import json
import math
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


def main(argv) -> int:
    from PySide6.QtCore import QEvent, Qt
    from PySide6.QtGui import QKeyEvent
    from PySide6.QtWidgets import QApplication

    from flowless import __main__ as entry
    from flowless import app as qapp
    from flowless import model

    assert "matplotlib" not in sys.modules, "the window module must not load matplotlib"
    from flowless import report
    args = [a for a in argv if not a.startswith("--")]
    shots = Path(argv[argv.index("--shots") + 1]) if "--shots" in argv else None
    if shots:
        args = [a for a in args if Path(a) != shots]
    data = Path(args[0]) if args else ROOT.parent / "data" / "data"
    run = data / "20260629_145214"
    out = Path(tempfile.mkdtemp(prefix="flowless_smoke_"))
    qapp.OPTIONS = out / "options.json"                  # leave the real options alone

    a = QApplication.instance() or QApplication([])
    entry._apply_dark(a)
    w = qapp.MainWindow()
    w.setAttribute(Qt.WA_DontShowOnScreen)
    w.resize(1500, 950)
    w.show()
    a.processEvents()

    def shot(name):
        a.processEvents()
        if shots:
            shots.mkdir(parents=True, exist_ok=True)
            w.grab().save(str(shots / f"{name}.png"))

    def key(widget, k):
        widget.keyPressEvent(QKeyEvent(QEvent.KeyPress, k, Qt.NoModifier))

    shot("0_empty")
    w.open_paths([run])
    exp = w.exp
    assert exp.groups() == ["OVCAR", "PC3"], exp.groups()
    assert all(exp.gates[g][0].name == model.AUTO for g in exp.groups()), "auto gate per group"
    from PySide6.QtWidgets import QDoubleSpinBox
    assert QDoubleSpinBox().textFromValue(1.5) == "1.50", "a point, whatever the system's locale"
    roles = {(s.group, s.label): s.role for s in exp.samples}
    assert roles[("PC3", "R3B23")] == "negative control", roles
    assert roles[("PC3", "AcII only")] == "secondary only"
    w.open_paths([data / "procedures" / "TheraSen_260629_FACS_autres_lignees.xlsx"])
    rec = w.exp.record
    assert len(rec["procedure"]) > 10 and any(r["catalog #"] == "AB9108" for r in rec["reagents"])
    assert any(r["kind"] == "sample" for r in rec["reagents"]), "the binders tested: kind sample"
    assert rec["groups"]["OVCAR"]["cell line"] == "OVCAR-3", rec["groups"]
    s = next(s for s in w.exp.samples if s.group == "PC3" and s.label == "29A1")
    assert "29A1" in s.fields.get("condition", "") and s.fields.get("secondary µL"), s.fields
    assert w.exp.fluorescence() == ["FL2-A"], "PE in the reagents: FL2-A only"
    assert w.exp.shown_channels() == ["FSC-A", "SSC-A", "FL2-A"], w.exp.shown_channels()
    w.samples_tab.table.selectRow(2)
    shot("1_samples")

    w.open_paths([run / "29-Jun-2026.wsp"])            # the workspace's gates replace the auto ones
    assert w.exp.gates_from().get("how") == "imported", w.exp.gates_file
    text, _ = report.build(w.exp, ["files"])
    assert "29-Jun-2026.wsp — the gates, imported" in text and "PC3_unstained" in text, text[-800:]
    vals = {(r["group"], r["label"]): r["value"] for r in report.summary_rows(w.exp)}
    assert abs(vals[("PC3", "29A1")] - 70543) / 70543 < 0.005, vals
    assert abs(vals[("PC3", "R3B23")] - 3854) / 3854 < 0.005, vals

    for i, name in enumerate(["1_samples", "2_gating", "3_figure", "4_results", "5_experiment"]):
        w.tabs.setCurrentIndex(i)
        if i:
            shot(name)
    w.tabs.setCurrentIndex(1)
    g = w.gating_tab
    from matplotlib.backend_bases import MouseEvent

    def mouse(kind, xy, canvas=g.dot, ax=g.ax1, px=None, dbl=False):
        x, y = px if px is not None else ax.transData.transform(xy)
        MouseEvent(kind, canvas, x, y, button=1 if kind != "motion_notify_event" else None,
                   dblclick=dbl)._process()

    def click(xy, **kw):
        mouse("motion_notify_event", xy, **kw)
        mouse("button_press_event", xy, **kw)
        mouse("button_release_event", xy, **kw)

    g.dot.draw()
    n0 = len(g.gates())
    # the dot plot's channels: chosen in a list (two usual pairs, the gates' pairs, custom)
    names = [g.axes.itemText(i) for i in range(g.axes.count())]
    assert names[-1] == "custom…" and any("FSC-H" in n for n in names), names
    assert not g.cx.isVisibleTo(g)
    g.axes.setCurrentIndex(g.axes.count() - 1)
    g._axes_picked(0)
    assert g.cx.isVisibleTo(g) and g.cx.count() == len(w.exp.all_channels())
    # the editor's markers stay out of the tight layout: the plot keeps its size (4.png)
    first = g.gates()[0]
    g.set_xy((first.x, first.y))
    assert not g.cx.isVisibleTo(g), "a listed pair: the custom lists hidden"
    g.pop = first.name
    g.fill_tree()
    g.draw()
    assert g.editor is not None, "the gate on its axes: editable"
    sizes = []
    g.dot.resize(700, 500)                             # the first resize: settled
    a.processEvents()
    g.dot.draw()
    for wd in (700, 500, 700, 500, 700):
        g.dot.resize(wd, 500)
        a.processEvents()
        g.dot.draw()
        sizes.append(tuple(round(v, 3) for v in g.ax1.get_position().bounds))
    assert sizes[1] == sizes[3] and sizes[2] == sizes[4], sizes   # no creeping in
    # the auto gate on all events: a new gate beside the others, not around them
    g.pop = model.ALL
    g.fill_tree()
    before = [(x.name, x.parent) for x in g.gates()]
    g.auto()
    after = [(x.name, x.parent) for x in w.exp.gates[g.group]]
    assert after[:-1] == before and after[-1] == (model.AUTO, ""), after
    w.do_undo()
    g = w.gating_tab
    # the histogram's y axis: count by default, % of events adds up to 100
    assert g.ny.currentData() == "count" and g.ax2.get_ylabel() == "Count", g.ax2.get_ylabel()
    g.ny.setCurrentIndex(g.ny.findData("area"))
    assert g.ax2.get_ylabel().startswith("% of events"), g.ax2.get_ylabel()
    ys = [ln.get_ydata() for ln in g.ax2.get_lines() if len(ln.get_ydata()) > 50]
    assert ys and all(abs(sum(y) - 100) < 0.5 for y in ys), [sum(y) for y in ys]
    g.ny.setCurrentIndex(g.ny.findData("count"))
    # the wheel over a list leaves it as it is
    from PySide6.QtCore import QPoint, QPointF
    from PySide6.QtGui import QWheelEvent
    i0 = g.ch.currentIndex()
    ev = QWheelEvent(QPointF(5, 5), QPointF(5, 5), QPoint(0, 0), QPoint(0, -120), Qt.NoButton,
                     Qt.NoModifier, Qt.NoScrollPhase, False)
    a.sendEvent(g.ch, ev)
    assert g.ch.currentIndex() == i0, "the wheel changed a list"
    # gates exported as Gating-ML 2.0, nested; the report names the file while they're unchanged
    from flowless.imports import gatingml
    import xml.etree.ElementTree as ET
    g.new_range()
    nx = gatingml.write(w.exp, out / "gates.xml")
    root = ET.parse(out / "gates.xml").getroot()
    G = "{http://www.isac-net.org/std/Gating-ML/v2.0/gating}"
    kids = [e for e in root if e.tag in (G + "PolygonGate", G + "RectangleGate")]
    assert nx == len(kids) == sum(len(v) for v in w.exp.gates.values()), (nx, len(kids))
    assert any(e.get(G + "parent_id") for e in kids), "nested gates keep their parent"
    poly = next(e for e in kids if e.tag == G + "PolygonGate")
    assert len(poly.findall(G + "vertex")) >= 3
    # and read back: the same gates, each in its group
    from dataclasses import asdict
    before = {k: [asdict(x) for x in v] for k, v in w.exp.gates.items()}
    e2, _ = model.Experiment.from_dict(w.exp.to_dict())
    for k in e2.gates:
        e2.gates[k] = []
    gatingml.apply(e2, out / "gates.xml")
    after = {k: [asdict(x) for x in v] for k, v in e2.gates.items()}
    assert after == before, (before, after)
    w.open_paths([out / "gates.xml"])                  # through the window: the file recorded
    assert {k: [asdict(x) for x in v] for k, v in w.exp.gates.items()} == before
    assert w.exp.gates_from().get("how") == "imported"
    # another program's file: no group (every group gets it); a transformed gate not read
    (out / "other.xml").write_text(f"""<?xml version="1.0"?>
<gating:Gating-ML xmlns:gating="{G[1:-1]}" xmlns:data-type="http://www.isac-net.org/std/Gating-ML/v2.0/datatypes">
 <gating:RectangleGate gating:id="Pos"><gating:dimension gating:min="5000">
  <data-type:fcs-dimension data-type:name="FL2-A"/></gating:dimension></gating:RectangleGate>
 <gating:RectangleGate gating:id="Logged"><gating:dimension gating:min="2" gating:transformation-ref="T1">
  <data-type:fcs-dimension data-type:name="FL2-A"/></gating:dimension></gating:RectangleGate>
</gating:Gating-ML>""", encoding="utf-8")
    from flowless import imports
    assert imports.kind(out / "other.xml") == "gatingml" and imports.kind(out / "gates.xml")
    (out / "not.xml").write_text("<root/>", encoding="utf-8")
    assert imports.kind(out / "not.xml") == ""
    gs_, skipped = gatingml.read(out / "other.xml")
    assert [x["name"] for x in gs_] == ["Pos"] and "Logged" in skipped[0], (gs_, skipped)
    w.open_paths([out / "other.xml"])
    assert all([x.name for x in w.exp.gates[k]] == ["Pos"] for k in w.exp.groups())
    w.do_undo()
    w.do_undo()
    assert {k: [asdict(x) for x in v] for k, v in w.exp.gates.items()} == before
    w.do_undo()
    g = w.gating_tab
    # a range under a polygon the figure doesn't use: said, and one click uses it
    gs = w.exp.gates[g.group]
    w.snap()
    gs.append(model.Gate("gate 9", "", "polygon", "FSC-A", "SSC-A",
                         [[1e5, 1e5], [1.6e7, 1e5], [1.6e7, 1.6e7]]))
    gs.append(model.Gate("hi 9", "gate 9", "range", "FL2-A", lo=1e4))
    w.changed()
    g = w.gating_tab
    assert g.warn.isVisibleTo(g) and "hi 9" in g.warn.text(), g.warn.text()
    assert not any(r["sub-population"] == "hi 9" for r in report.summary_rows(w.exp))
    g._use(str(len(gs) - 2))
    from flowless import plots
    assert plots.measured(w.exp, g.group) == "gate 9" and not g.warn.isVisibleTo(g)
    sr = [r for r in report.summary_rows(w.exp) if r["group"] == g.group]
    assert any(r["sub-population"] == "hi 9" for r in sr), "the summary follows the figure"
    assert any(r["role"] == "negative control" and r["÷ control"] == 1.0
               for r in sr if not r["sub-population"]), sr
    w.do_undo()
    w.do_undo()
    g = w.gating_tab
    g.set_xy(("FSC-A", "SSC-A"))
    g.dot.draw()
    n0 = len(g.gates())
    g.new_polygon()
    g.dot.draw()
    corners = [(3e6, 0.5e6), (8e6, 0.5e6), (8e6, 3e6), (3e6, 3e6)]
    for xy in corners + corners[:1]:
        click(xy)
    a.processEvents()
    assert len(g.gates()) == n0 + 1, "a polygon drawn with the mouse becomes a gate"
    new = g.gate()
    assert new is not None and len(new.vertices) == 4, new
    g.dot.draw()
    # drag a corner
    mouse("motion_notify_event", (8e6, 3e6))
    mouse("button_press_event", (8e6, 3e6))
    mouse("motion_notify_event", (10e6, 4e6))
    mouse("button_release_event", (10e6, 4e6))
    assert any(abs(x - 10e6) < 3e5 and abs(y - 4e6) < 3e5 for x, y in g.gate().vertices), \
        g.gate().vertices
    # a click on a side adds a corner there; a corner selected then Del deletes it
    g.dot.draw()
    click((5.5e6, 0.5e6))
    assert len(g.gate().vertices) == 5, g.gate().vertices
    g.dot.draw()
    click((5.5e6, 0.5e6))
    key(g.dot, Qt.Key_Delete)
    assert len(g.gate().vertices) == 4, g.gate().vertices
    # drawing: a click outside the plot puts the corner on its edge (the top of the scale)
    g.new_polygon()
    g.dot.draw()
    top = g.samples()[0].fcs.range("FSC-A")
    for xy in [(12e6, 0.5e6), (14e6, 0.5e6)]:
        click(xy)
    bx = g.ax1.bbox
    click(None, px=(bx.x1 + 15, g.ax1.transData.transform((0, 2e6))[1]))
    click((12e6, 2.5e6))
    key(g.dot, Qt.Key_Return)
    edge = g.gate()
    assert edge is not None and any(abs(x - top) < 1 for x, _ in edge.vertices), edge
    shot("2_gating_drawn")
    for _ in range(5):
        w.do_undo()
    assert len(w.exp.gates[g.group]) == n0, "undo every step"
    g = w.gating_tab
    g.new_range()
    assert any(x.kind == "range" for x in w.exp.gates[g.group]), "range gate added"
    shot("2_gating_range")
    w.do_undo()
    assert not any(x.kind == "range" for x in w.exp.gates[g.group]), "undo"
    w.do_redo()
    assert any(x.kind == "range" for x in w.exp.gates[g.group]), "redo"
    w.do_undo()
    # a sample's own gate: it alone changes
    g = w.gating_tab
    # Groups / samples: several samples of a group, then one; the axes stay as set
    from PySide6.QtCore import QItemSelectionModel
    t = g.groups
    top = next(t.topLevelItem(i) for i in range(t.topLevelItemCount())
               if t.topLevelItem(i).data(0, Qt.UserRole)[1] == g.group)
    g.set_xy(("FSC-A", "FSC-H"))
    t.clearSelection()
    for i in (0, 1):
        top.child(i).setSelected(True)
    t.setCurrentItem(top.child(1), 0, QItemSelectionModel.NoUpdate)
    a.processEvents()
    assert len(g.samples()) == 2 and g.one() is None and not g.own.isEnabled(), g.picked
    two = sum(int(w.exp.masks(s)[model.ALL].sum()) for s in g.samples())
    assert g.tree.topLevelItem(0).text(2) == str(two), "the counts: those two samples"
    g.pop = g.gates()[0].name
    g.fill_tree()
    g._pop_changed()
    assert g.xy == ("FSC-A", "FSC-H"), "selecting a population keeps the plot's axes"
    t.clearSelection()
    top.child(0).setSelected(True)
    t.setCurrentItem(top.child(0), 0, QItemSelectionModel.NoUpdate)
    a.processEvents()
    assert g.one() is not None and len(g.samples()) == 1
    g.set_xy((g.gates()[0].x, g.gates()[0].y))
    first = g.gates()[0]
    g.pop = first.name
    g.fill_tree()
    g.own.setChecked(True)
    one = g.one()
    assert first.name in one.gates, "own gate made"
    g._poly_done([[1e6, 1e5], [1.6e7, 1e5], [1.6e7, 1.6e7], [1e6, 1.6e7]])
    assert w.exp.gates[g.group][0].vertices != one.gates[first.name].vertices
    r = {x["sample"]: x["own gate"] for x in w.exp.results() if x["population"] == first.name}
    assert r[one.name] and sum(r.values()) == 1, r
    g.own.setChecked(False)
    assert first.name not in g.one().gates

    # the figure: panels, options, the preview
    from flowless import plots
    w.tabs.setCurrentIndex(2)
    f = w.figure_tab
    n = len(f.panels())
    f.items.setCurrentRow(1)
    f.add()
    f._panel_set("kind", "dot")
    f.items.setCurrentRow(0)
    f.preset.setCurrentText("black & white, offset")
    f._apply_preset()
    a.processEvents()
    f.render()
    assert len(f.panels()) == n + 1 and f.panels()[1]["kind"] == "dot", f.panels()
    x0, y0, x1, y1 = dict(f.fig.flowless_cells)[1]
    f.view.clicked.emit((x0 + x1) / 2, (y0 + y1) / 2)
    assert f.sel == 1, f.sel
    f._set("lw", 2.0, f.panels()[f.sel])
    assert f.panels()[1]["opts"]["lw"] == 2.0, "a panel's own option"
    from PySide6.QtWidgets import QDoubleSpinBox, QSpinBox
    spins = f.opt_box.findChildren(QDoubleSpinBox)
    assert spins and all("," not in x.text() for x in spins), [x.text() for x in spins]
    from flowless.figtab import Num
    n = Num(singleStep=0.1)
    assert n.valueFromText("0.75") == 0.75 and n.textFromValue(1.1) == "1.1" and \
        n.textFromValue(0.75) == "0.75"
    f._set("axis_lw", 1.5, None)
    f.render()
    assert any(abs(sp.get_linewidth() - 1.5) < 1e-6 for ax in f.fig.axes for sp in ax.spines.values())
    # the layout, with the mouse: + adds a panel (3 + 1), one dragged down makes 2 × 2 and
    # nothing else moves; over another box: not taken; half a panel down; too wide: red, not
    # drawn; the legend into the grid, then above the figure; the list's order moves nothing
    from PySide6.QtCore import QPoint
    from PySide6.QtTest import QTest
    sc = f.schema
    a.processEvents()

    Q = plots.Q

    def at():                                  # in halves of a panel, as first written
        return {i: (v[0] * 2 / Q, v[1] * 2 / Q) for i, v in f.fig.flowless_grid["at"].items()}

    def drag(src, dst):
        src = src.toPoint() if hasattr(src, "toPoint") else src
        QTest.mousePress(sc, Qt.LeftButton, Qt.NoModifier, src)
        QTest.mouseMove(sc, src + QPoint(12, 12))
        QTest.mouseMove(sc, dst)
        QTest.mouseRelease(sc, Qt.LeftButton, Qt.NoModifier, dst)
        a.processEvents()

    def spot(i, x, y):                         # the pointer that puts item i at (x, y) halves
        w_, h_ = sc._size(i)
        return sc._px(*sc._q(x * Q / 2 + w_ / 2, y * Q / 2 + h_ / 2, 0, 0)).topLeft().toPoint()

    assert len(sc._boxes()) == len(f.panels()), (sc._boxes(), f.panels())
    assert sorted(at().values()) == [(0, 0), (2, 0), (4, 0)], at()
    QTest.mouseClick(sc, Qt.LeftButton, Qt.NoModifier, sc._px(*sc._ghost()).center().toPoint())
    assert sorted(at().values()) == [(0, 0), (0, 2), (2, 0), (4, 0)], at()
    third = next(i for i, v in at().items() if v == (4, 0))
    drag(sc._boxes()[third].center(), spot(third, 2, 2))
    assert sorted(at().values()) == [(0, 0), (0, 2), (2, 0), (2, 2)], at()
    keep = at()
    drag(sc._boxes()[third].center(), spot(third, 1, 1))
    assert at() == keep, "over the others: not taken"
    drag(sc._boxes()[third].center(), spot(third, 4, 1))
    assert at()[third] == (4, 1) and {i: v for i, v in at().items() if i != third} == \
        {i: v for i, v in keep.items() if i != third}, at()
    cells = dict(f.fig.flowless_cells)
    H = f.fig.get_figheight() * 25.4
    assert abs((1 - cells[third][3]) * H - f.fig.flowless_grid["cell"][1] / 2) < 0.01, \
        "half a panel down"
    assert not f.wide.isVisible() and not f.too_wide()
    cw = f.fig.flowless_grid["cell"][0]
    far = 2 * int(plots.options(w.exp)["max_w"] // cw)
    drag(sc._boxes()[third].center(), spot(third, far, 1))
    assert f.too_wide(), (at(), far, f.fig.flowless_grid["wide"], plots.figure(w.exp, draw=False))
    assert not f.wide.isHidden(), f.wide.text()
    assert not f.fig.axes and f.max_w_box.styleSheet() and f._refuse(), "not drawn"
    shot("3_figure_wide")
    w.do_undo()
    f.refresh()
    assert at()[third] == (4, 1) and not f.too_wide() and f.wide.isHidden()
    li = sc.legend_i
    drag(sc._boxes()[li].center(), spot(li, 0, 4))
    assert plots.options(w.exp)["legend"] == "grid" and at()[li] == (0, 4), at()
    drag(sc._boxes()[li].center(), sc._zones()["top"].center().toPoint())
    assert plots.options(w.exp)["legend"] == "top" and li not in at(), at()
    before = {id(q): tuple(q["pos"]) for q in f.panels() if q.get("kind") != "legend"}
    f.items.insertItem(f.items.count(), f.items.takeItem(1))
    f._moved()
    assert {id(q): tuple(q["pos"]) for q in f.panels() if q.get("kind") != "legend"} == before
    f.schema.picked.emit(1)
    assert f.sel == 1 and f.items.currentRow() == 2, "the box picks the list's line"
    f.build_bar()
    cols = next(x for x in f.bar_box.findChildren(QSpinBox)
                if x.specialValueText() == "as many as fit")
    cols.setValue(1)
    f.render()
    assert {v[0] for v in at().values()} == {0}, "one column, the places reset"
    cols.setValue(0)
    f.render()
    W0 = plots.figure(w.exp, draw=False)[0]
    f.res.setValue(20)
    f.render()
    assert f.size.text().startswith(f"{2 * W0:.0f}"), (f.size.text(), W0)
    f.res.setValue(10)
    # titles keep their size (a long one wraps), a legend in the grid sits where set, x
    # ticks thinned to what fits, the ranges' labels apart
    for k, v in (("title_pt", 15.0), ("xtick_pt", 12.0), ("ytick_pt", 12.0), ("panel_w", 40.0),
                 ("legend", "grid"),
                 ("legend_h", "left"), ("legend_dx", 5.0)):
        f._set(k, v, None)
    w.set_alias("group", w.exp.groups()[0], "a rather long name for this group",
                remember=False)
    gr = w.exp.groups()[0]
    rng = w.exp.gates[gr]
    pop = w.exp.main_population(gr)
    rng += [model.Gate("a long range name", pop, "range", "FL2-A", lo=1e3, hi=1.5e3),
            model.Gate("another long one", pop, "range", "FL2-A", lo=1.5e3, hi=2e3)]
    fig = plots.figure(w.exp)
    fig.canvas.draw()
    titles = [ax.title for ax in fig.axes if ax.get_title()]
    assert titles and all(t.get_fontsize() == 15.0 for t in titles), \
        [(t.get_text(), t.get_fontsize()) for t in titles]
    assert any("\n" in t.get_text() for t in titles), [t.get_text() for t in titles]
    leg = fig.legends[0].get_window_extent()
    li = next(i for i, p in enumerate(plots.panels(w.exp)) if p.get("kind") == "legend")
    cx0 = dict(fig.flowless_cells)[li][0] * fig.bbox.width
    assert abs(leg.x0 - cx0 - 5 / 25.4 * fig.dpi) < 0.6 / 25.4 * fig.dpi, (leg.x0, cx0)
    hax = next(ax for ax in fig.axes if ax.get_xscale() == "log")
    shown = [t for t in hax.get_xticklabels() if t.get_text() and
             hax.get_xlim()[0] <= t.get_position()[0] <= hax.get_xlim()[1]]
    assert len(shown) < 8, [t.get_text() for t in shown]
    labels = [t for ax in fig.axes for t in ax.texts if t.get_text() in
              ("a long range name", "another long one")]
    assert len(labels) == 2 and labels[0].get_position()[1] != labels[1].get_position()[1], \
        "two labels wider than their ranges: one row each"
    assert all(t.get_position()[1] > 1 for t in labels), "the ranges above the plot"

    def mm(ax):
        return (round(ax.get_position().width * fig.get_figwidth() * 25.4, 2),
                round(ax.get_position().height * fig.get_figheight() * 25.4, 2))

    assert len({mm(ax) for ax in fig.axes}) == 1, [mm(ax) for ax in fig.axes]
    del rng[-2:]
    w.do_undo()
    for _ in range(7):
        w.do_undo()
    w.set_alias("group", gr, "", remember=False)
    W0, H0 = plots.figure(w.exp, draw=False)
    f._set("size_of", "plot", None)            # the size set is the plot's, the figure stays
    fig = plots.figure(w.exp)
    o = plots.options(w.exp)
    assert {mm(ax) for ax in fig.axes} == {(round(o["panel_w"], 2), round(o["panel_h"], 2))}, \
        ([mm(ax) for ax in fig.axes], o["panel_w"], o["panel_h"])
    W1, H1 = plots.figure(w.exp, draw=False)
    assert abs(W1 - W0) < 0.1 and abs(H1 - H0) < 0.5, (W0, H0, W1, H1)
    w.do_undo()
    assert plots.options(w.exp)["size_of"] == "panel"
    # 9b.png: the legend on 2 lines in half a place, a panel half a place below it
    keep = json.loads(json.dumps(w.exp.figure))
    w.exp.figure["options"].update(legend="grid", legend_rows=2)
    f.render()
    lw, lh = f.fig.flowless_grid["legend"]
    own = f._own_panels()
    leg = next(q for q in own if q.get("kind") == "legend")
    pan = [q for q in own if q.get("kind") != "legend"]
    Q = plots.Q
    leg["pos"], pan[0]["pos"] = [0, 0], [0, 0.5]
    x0 = max(1, math.ceil(lw / Q))
    for k, q in enumerate(pan[1:]):
        q["pos"] = [x0 + k % 2, k // 2]
    w.exp.figure["options"]["max_w"] = 800.0
    fig = plots.figure(w.exp)
    cells, ch = dict(fig.flowless_cells), fig.flowless_grid["cell"][1]
    H = fig.get_figheight() * 25.4
    li, pi = own.index(leg), own.index(pan[0])
    assert lh <= Q // 2 and abs((cells[li][3] - cells[li][1]) * H - lh * ch / Q) < 0.01, \
        (lh, cells[li])
    assert abs((1 - cells[pi][3]) * H - ch / 2) < 0.01, "the panel under the legend"
    assert len({mm(ax) for ax in fig.axes}) == 1, [mm(ax) for ax in fig.axes]
    n = len(fig.legends[0].get_texts())
    assert fig.legends[0]._ncols == -(-n // 2), (n, fig.legends[0]._ncols)
    f.refresh()
    shot("3_figure_panels")
    # the legend's handle: a quarter of a panel high, then back to its entries' size
    hd = f.schema._handle().center().toPoint()
    top_left = f.schema._boxes()[li].topLeft().toPoint()
    to = f.schema._px(*f.schema._q(0, 0, lw, 1)).bottomRight().toPoint()
    QTest.mousePress(f.schema, Qt.LeftButton, Qt.NoModifier, hd)
    QTest.mouseMove(f.schema, to)
    QTest.mouseRelease(f.schema, Qt.LeftButton, Qt.NoModifier, to)
    a.processEvents()
    own = f.panels()
    assert own[li].get("size") == [lw / Q, 1 / Q], own[li].get("size")
    assert f.fig.flowless_grid["at"][li][2:] == (lw, 1), f.fig.flowless_grid["at"][li]
    f._legend_size(li, None, None)
    assert "size" not in f.panels()[li]
    # a panel's own shift (mm): its plot moves, the others don't; dragging it resets it
    f.items.setCurrentRow(pi + 1)
    ax0 = [tuple(ax.get_position().bounds) for ax in f.fig.axes]
    f._shift(0, 3.0)
    f._shift(1, 2.0)
    f.render()
    ax1 = [tuple(ax.get_position().bounds) for ax in f.fig.axes]
    moved = [k for k, (u, v) in enumerate(zip(ax0, ax1)) if u != v]
    assert len(moved) == 1, moved
    k = moved[0]
    W, H = f.fig.get_figwidth() * 25.4, f.fig.get_figheight() * 25.4
    assert abs((ax1[k][0] - ax0[k][0]) * W - 3) < 0.01 and abs((ax0[k][1] - ax1[k][1]) * H - 2) \
        < 0.01, (ax0[k], ax1[k])
    f._place(pi, 0, 2, "")
    assert "shift" not in f.panels()[pi]
    # the legend in each panel: aligned in the plot as set (right, top, 3 mm in)
    w.exp.figure["options"].update(legend="panel", legend_h="right", legend_v="top",
                                   legend_dx=3.0, legend_dy=3.0)
    fig = plots.figure(w.exp)
    fig.canvas.draw()
    ax = next(x for x in fig.axes if x.get_legend())
    lb, ab = ax.get_legend().get_window_extent(), ax.get_window_extent()
    px = fig.dpi / 25.4
    assert abs(ab.x1 - lb.x1 - 3 * px) < 0.5 * px and abs(ab.y1 - lb.y1 - 3 * px) < 0.5 * px, \
        (lb, ab)
    w.exp.figure.clear()
    w.exp.figure.update(keep)
    f.refresh()
    f.items.setCurrentRow(0)
    f.preset.setCurrentText("colour")
    f._apply_preset()
    # the series' order: ▲ on the last series puts it before the one above, everywhere
    f.items.setCurrentRow(0)
    o = plots.options(w.exp)
    names = list(f._series_names())
    assert len(names) >= 2, names
    f._series_move(len(names) - 1, -1)
    want = names[:-2] + [names[-1], names[-2]]
    assert list(f._series_names()) == want, (names, list(f._series_names()))
    leg = plots._legend_entries(w.exp, plots.panels(w.exp), plots.series_style(w.exp, o), o)[1]
    assert [n for n in leg if n in want] == [n for n in want if n in leg], leg
    rk = w.exp.rank()
    sr = report.summary_rows(w.exp)
    for g in w.exp.groups():
        got = [rk[r["label"]] for r in sr if r["group"] == g]
        assert got == sorted(got), (g, got)
    assert model.Experiment.from_dict(w.exp.to_dict())[0].labels() == w.exp.labels()
    w.do_undo()
    assert list(f._series_names()) == names and not w.exp.label_order
    # Results: sorted by a column's title, the other way, then as computed
    rt = w.results_tab
    rt.mode.setCurrentIndex(0)
    rt.refresh()
    first = [rt.table.item(r, 0).text() for r in range(rt.table.rowCount())]
    k = [rt.table.horizontalHeaderItem(c).text() for c in range(rt.table.columnCount())] \
        .index(w.exp.stat)
    rt._sort(k)
    vals = [rt.rows()[1][r][k] for r in range(rt.table.rowCount())]
    assert vals == sorted(vals) and rt.table.horizontalHeaderItem(k).text().endswith("▲")
    rt._sort(k)
    assert [rt.rows()[1][r][k] for r in range(len(vals))] == sorted(vals, reverse=True)
    rt.table.selectRow(0)
    rt.copy()
    assert str(rt.rows()[1][0][0]) in QApplication.clipboard().text().splitlines()[1]
    rt._sort(k)
    assert [rt.table.item(r, 0).text() for r in range(rt.table.rowCount())] == first
    assert rt.table.verticalHeader().isHidden() and "—" not in \
        [rt.table.item(r, 2).text() for r in range(rt.table.rowCount())], "the population named"
    for ext in (".png", ".svg", ".pdf", ".tif"):
        f._write(out / f"fig{ext}")
        assert (out / f"fig{ext}").stat().st_size > 1000, ext
    f.copy()
    # aliases
    w.set_alias("channel", "FL2-A", "anti-His PE", remember=False)
    assert w.exp.name("channel", "FL2-A") == "anti-His PE"
    w.changed()
    report.save_xlsx(w.exp, out / "values.xlsx")
    text, images = report.build(w.exp, None, True)
    assert "anti-His PE" in text and "PE-labelled anti-His-tag secondary antibody" in text
    report.save_html(text, images, out / "report.html")
    qapp.to_pdf(text, images, out / "report.pdf")
    assert (out / "report.pdf").stat().st_size > 20000
    # a singlets gate: first under the main population, so the figure measures it
    w.tabs.setCurrentIndex(1)
    g = w.gating_tab
    g.pop = w.exp.main_population(g.group)
    g.fill_tree()
    main = g.pop
    g.singlets()
    assert w.exp.main_population(g.group) == model.SINGLETS, w.exp.gates[g.group]
    sg = next(x for x in w.exp.gates[g.group] if x.name == model.SINGLETS)
    assert (sg.parent, sg.x, sg.y) == (main, "FSC-A", "FSC-H") and g.xy == ("FSC-A", "FSC-H")
    n_in = [r["% parent"] for r in w.exp.results() if r["population"] == model.SINGLETS]
    assert all(45 < x < 90 for x in n_in), n_in   # OVCAR's gate on FSC-H × SSC-H: ~40 % doublets
    w.do_undo()
    # two runs of one cell line, their folders named alike: pooled as replicates
    w.open_paths([data / "20260527_165837" / "MDA-MB-157 parental cells",
                  data / "20260611_152206" / "MDA-MB-157-parental cells"])
    pooled = "MDA-MB-157-parental cells"
    assert w.exp.name("group", pooled) == "MDA-MB-157 parental cells", w.exp.aliases
    assert w.exp.title(pooled).endswith("2026-06-11"), w.exp.title(pooled)
    rr = report.replicate_rows(w.exp)
    assert rr and all(r["n"] == 2 and r["sd"] is not None for r in rr), rr
    ctrl = next(r for r in rr if r["label"] == w.exp.name("label", "0006"))
    assert abs(ctrl["ratio mean"] - 1) < 1e-9, ctrl
    w.tabs.setCurrentIndex(3)
    w.results_tab.mode.setCurrentIndex(2)
    assert w.results_tab.table.rowCount() == len(rr)
    w.results_tab.mode.setCurrentIndex(0)
    f = w.figure_tab
    own = f._own_panels()
    bars = plots.new_panel(pooled, "bars")
    own.insert(0, bars)
    fig = plots.figure(w.exp)
    ax = fig.axes[0]
    assert "2 runs" in ax.get_title() and ax.containers, "the pooled bars: mean ± SD"
    own.remove(bars)
    # the run's day in a title only while another run of it has a panel too
    hp = next(p for p in own if p.get("group") == pooled)
    first = next(p for p in own if p.get("group") == "MDA-MB-157 parental cells")
    assert plots.panel_title(w.exp, hp).endswith("2026-06-11"), plots.panel_title(w.exp, hp)
    k = own.index(first)
    own.remove(first)
    assert plots.panel_title(w.exp, hp) == "MDA-MB-157 parental cells"
    own.insert(k, first)
    # labels shown under one name: one series (one style, one legend entry)
    o = plots.options(w.exp)
    labs = [lab for lab in w.exp.labels() if any(s.label == lab and s.plot
                                                  for s in w.exp.samples)][:2]
    for lab in labs:
        w.set_alias("label", lab, "one name", remember=False)
    st = plots.series_style(w.exp, o)
    assert st[labs[0]] is st[labs[1]]
    names = plots._legend_entries(w.exp, plots.panels(w.exp), st, o)[1]
    assert names.count("one name") == 1, names
    for lab in labs:
        w.set_alias("label", lab, "", remember=False)
    # the report: each format its own sections; a PowerPoint slide, an Excel by section
    rd = qapp.ReportDialog(w)
    rd.fmt.setCurrentIndex(rd.fmt.findData(".pdf"))
    assert "SHA-256" in rd.view.toPlainText() and rd.list.count() == len(report.SECTIONS)
    rd.list.item(0).setCheckState(Qt.Unchecked)
    rd.refresh()
    assert w.opts["report_sections"][".pdf"]["off"] == [rd.list.item(0).data(Qt.UserRole)]
    rd.fmt.setCurrentIndex(rd.fmt.findData(".pptx"))
    ticked = rd.sections()
    assert "samples" not in ticked and "experiment" in ticked, ticked
    assert "slide" in rd.list.item(1).text() or "notes" in rd.list.item(1).text()
    assert "Notes (below the slide)" in rd.view.toPlainText()
    rd.write(out / "report.pptx")
    from pptx import Presentation
    deck = Presentation(out / "report.pptx")
    assert len(deck.slides) == 1
    shapes = {sh.name for sh in deck.slides[0].shapes}
    assert {"Title", "Figure"} <= shapes and any(n in shapes for n in ("Summary",)), shapes
    assert any(sh.has_table and len(sh.table.rows) > 5 for sh in deck.slides[0].shapes)
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    assert "PROCEDURE AND REAGENTS" in notes and "AB9108" in notes, notes[:400]
    rd.fmt.setCurrentIndex(rd.fmt.findData(".xlsx"))
    rd.write(out / "report.xlsx")
    from openpyxl import load_workbook
    sheets = load_workbook(out / "report.xlsx").sheetnames
    assert "Summary" in sheets and "Replicates" in sheets and "Figure" not in sheets, sheets
    rd.fmt.setCurrentIndex(rd.fmt.findData(".pdf"))
    assert rd.list.item(0).checkState() == Qt.Unchecked, "the PDF's own sections came back"
    if shots:
        rd.fmt.setCurrentIndex(rd.fmt.findData(".pptx"))
        rd.resize(1300, 850)
        rd.grab().save(str(shots / "6_report.png"))
    rd.close()
    od = qapp.OptionsDialog(w)
    assert od.topics.count() == 3
    if shots:
        od.topics.setCurrentRow(2)
        od.grab().save(str(shots / "7_options.png"))
    od.close()
    ad = qapp.AliasDialog(w)                   # Ctrl+C a name, Ctrl+V as two rows' alias
    ad.table.setCurrentCell(1, 1)
    ad._copy()
    from PySide6.QtWidgets import QTableWidgetSelectionRange
    ad.table.clearSelection()
    ad.table.setRangeSelected(QTableWidgetSelectionRange(2, 0, 3, 2), True)
    ad._paste()
    want = ad.table.item(1, 1).text()
    assert [ad.table.item(r, 2).text() for r in (2, 3)] == [want, want]
    ad.close()
    w.tabs.setCurrentIndex(0)
    w.samples_tab._sort(w.samples_tab.shown().index("label"))
    assert [s.label for s in w.exp.samples] == sorted([s.label for s in w.exp.samples],
                                                      key=lambda x: model._natural(Path(x)))

    w.exp.save(out / "exp.json")
    d = json.loads((out / "exp.json").read_text(encoding="utf-8"))
    assert d["samples"][0]["rel"], "relative paths kept"
    e2, problems = model.Experiment.open(out / "exp.json")
    assert not problems and len(e2.samples) == len(w.exp.samples), problems
    assert e2.results() == w.exp.results() and e2.aliases == w.exp.aliases
    w.dirty = False
    w.close()
    print(f"smoke test OK — outputs in {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
